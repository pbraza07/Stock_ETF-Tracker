"""Exports and summary presentation from immutable YTD ledgers only."""
from io import BytesIO
from html import escape
import json
import pandas as pd
from reportlab.pdfgen import canvas
from reportlab.platypus import Table, Paragraph
from marketscope_pdf_theme import (PAGE_SIZE, BACKGROUND, HEADER_WASH, CARD, BORDER,
    ACCENT, CYAN, TEXT, MUTED, build_styles, table_style, marketscope_version)


def coverage_label(payload):
    if payload.get('custom_start'):
        return 'Selected period' if payload['full_ytd'] else 'Partial selected period'
    return 'Full YTD' if payload['full_ytd'] else 'Partial YTD'


def metrics(payload):
    s = dict(payload['snapshot'])
    s['Profit / Loss'] = s['Current Balance'] + s['Withdrawn'] - s['Beginning Balance']
    s['Total Return %'] = s['Profit / Loss'] / s['Beginning Balance'] * 100
    return s


def presentation_record(record):
    """Display existing None-cadence saves as one performance series as well."""
    if record.get('inputs',{}).get('cadence') == 'None':
        record=dict(record)
        record['strategies']={'Performance':next(iter(record['strategies'].values()))}
    return record


def performance_table(table):
    return table.drop(columns=['Requested Withdrawal','Actual Withdrawal','Shortfall','Cumulative Withdrawals'],errors='ignore').rename(columns={'Net Profit incl. Withdrawals':'Cumulative Profit'})


def summary_html(record):
    """Reuse the historical Portfolio card's exact classes and responsive layout."""
    record=presentation_record(record)
    performance='Performance' in record['strategies']
    first=next(iter(record['strategies'].values()))
    s=metrics(first)
    inputs=record['inputs']
    def sign(value):
        return 'pos' if value>0 else 'neg' if value<0 else 'flat'
    def field(label,value,number=None,secondary=False):
        cls='simulation-library-withdrawal-metric' if secondary else 'simulation-library-metric'
        color=sign(number) if number is not None else ''
        return f"<div class='{cls}'><small>{escape(label)}</small><b class='{color}'>{escape(value)}</b></div>"
    meta=f"{coverage_label(first)} · {first['start_date']} to {first['through']} · {len(inputs.get('holdings',[]))} instrument(s)"
    identity="<div class='simulation-library-identity'><span class='simulation-library-name'>"+escape(record.get('name','Simulation'))+"</span><small>"+escape(meta)+"</small></div>"
    main=[field('INVESTED',f"${s['Beginning Balance']:,.2f}"),
          field('ENDING' if performance else 'RB ENDING',f"${s['Current Balance']:,.2f}"),
          field('PROFIT / LOSS' if performance else 'RB PROFIT / LOSS',f"${s['Profit / Loss']:+,.2f}",s['Profit / Loss']),
          field('RETURN' if performance else 'RB RETURN',f"{s['Total Return %']:+.2f}%",s['Total Return %'])]
    if performance:
        lower=[field('POSITIVE DAYS',f"{s['Positive Days']}/{s['Days']}",secondary=True),
               field('POSITIVE MONTHS',f"{s['Positive Months']}/{s['Months']}",secondary=True)]
    else:
        nr=metrics(record['strategies']['Non-Rebalanced'])
        diff=s['Current Balance']-nr['Current Balance']
        lower=[field(inputs.get('cadence','Monthly').upper()+' WITHDRAWAL',f"${inputs.get('withdrawal',0):,.2f}",secondary=True),
               field('NOT-REBALANCED REMAINING',f"${nr['Current Balance']:,.2f}",secondary=True),
               field('NR PROFIT / RETURN',f"${nr['Profit / Loss']:+,.2f} / {nr['Total Return %']:+.2f}%",nr['Profit / Loss'],True),
               field('REBALANCE DIFFERENCE',f"${diff:+,.2f}",diff,True),
               field('POSITIVE DAYS',f"RB {s['Positive Days']}/{s['Days']} · NR {nr['Positive Days']}/{nr['Days']}",secondary=True),
               field('POSITIVE MONTHS',f"RB {s['Positive Months']}/{s['Months']} · NR {nr['Positive Months']}/{nr['Months']}",secondary=True)]
    return "<div class='simulation-library-card'>"+identity+''.join(main)+"<div class='simulation-library-withdrawal-strip'>"+''.join(lower)+"</div></div>"


def build_ytd_pdf(record):
    record=presentation_record(record)
    out = BytesIO()
    c = canvas.Canvas(out, pagesize=PAGE_SIZE)
    w,h = PAGE_SIZE
    styles = build_styles()
    page = 0
    def begin(title, subtitle):
        nonlocal page
        if page: c.showPage()
        page += 1
        c.setFillColor(BACKGROUND); c.rect(0,0,w,h,fill=1,stroke=0)
        c.setFillColor(HEADER_WASH); c.rect(0,h-112,w,112,fill=1,stroke=0)
        c.setFillColor(ACCENT); c.setFont('Helvetica-Bold',20)
        c.drawCentredString(w/2,h-36,title)
        c.setFillColor(MUTED); c.setFont('Helvetica',8)
        c.drawCentredString(w/2,h-55,subtitle[:150])
        c.setFont('Helvetica',7)
        c.drawRightString(w-26,h-16,'MarketScope v'+marketscope_version())
        c.drawString(26,18,'Historical adjusted-close simulation. Positive months include partial months. Not a forecast.')
        c.drawRightString(w-26,18,f'Page {page}')
    def cards(items, top, height=68):
        gap=10; cw=(w-52-gap*(len(items)-1))/len(items)
        for i,(label,value) in enumerate(items):
            x=26+i*(cw+gap)
            c.setFillColor(CARD); c.setStrokeColor(BORDER)
            c.roundRect(x,top-height,cw,height,7,fill=1,stroke=1)
            c.setFillColor(MUTED); c.setFont('Helvetica-Bold',7)
            c.drawCentredString(x+cw/2,top-18,label)
            c.setFillColor(ACCENT); c.setFont('Helvetica-Bold',12)
            c.drawCentredString(x+cw/2,top-40,value)
    inputs=record['inputs']
    for name,payload in record['strategies'].items():
        s=metrics(payload)
        mode=coverage_label(payload)
        subtitle=f"{record['name']} | {name} | {mode} {payload['start_date']} to {payload['through']} | Equal split"
        begin('PORTFOLIO SPLIT SIMULATOR',subtitle)
        cards([('TOTAL INVESTED',f"${s['Beginning Balance']:,.2f}"),('CURRENT BALANCE',f"${s['Current Balance']:,.2f}"),
               ('PROFIT / LOSS',f"${s['Profit / Loss']:+,.2f}"),('TOTAL RETURN',f"{s['Total Return %']:+.2f}%")],h-70)
        cards([('TOTAL WITHDRAWN',f"${s['Withdrawn']:,.2f}"),('POSITIVE DAYS',f"{s['Positive Days']}/{s['Days']}"),
               ('POSITIVE MONTHS',f"{s['Positive Months']}/{s['Months']}"),('WITHDRAWAL / '+inputs.get('cadence','Monthly').upper(),f"${inputs.get('withdrawal',0):,.2f}")],h-153,60)
        c.setFillColor(MUTED); c.setFont('Helvetica',8)
        c.drawString(26,h-230,'Profit includes withdrawals. Return = profit / initial investment; positive periods use returns before withdrawals.')
        c.setFillColor(ACCENT); c.setFont('Helvetica-Bold',11)
        c.drawString(26,h-258,'CUMULATIVE PROFIT / LOSS')
        daily=pd.DataFrame(payload['daily'])
        vals=[0.0]+(daily['Ending Balance']+daily['Actual Withdrawal'].cumsum()-s['Beginning Balance']).tolist()
        lo,hi=min(vals),max(vals)
        if lo==hi: lo-=1; hi+=1
        x,y,cw,ch=90,90,w-130,205
        c.setStrokeColor(BORDER); c.rect(x,y,cw,ch)
        c.setFillColor(MUTED); c.setFont('Helvetica',8)
        for j in range(5):
            value=lo+(hi-lo)*j/4
            yy=y+ch*j/4
            c.drawRightString(x-8,yy,f'${value:,.0f}')
            c.setStrokeColor(BORDER); c.line(x,yy,x+cw,yy)
        p=c.beginPath()
        for j,value in enumerate(vals):
            xx=x+cw*j/max(1,len(vals)-1); yy=y+ch*(value-lo)/(hi-lo)
            if j==0:p.moveTo(xx,yy)
            else:p.lineTo(xx,yy)
        c.setStrokeColor(CYAN); c.setLineWidth(1.5); c.drawPath(p)
        dates=[payload['start_date']]+daily['Date'].astype(str).str[:10].tolist()
        for j in sorted(set(round(k*(len(dates)-1)/4) for k in range(5))):
            c.drawCentredString(x+cw*j/max(1,len(dates)-1),y-16,dates[j])
        # A separate snapshot page preserves the reference report's instrument section.
        instruments=record.get('instruments') or [{'Symbol':s} for s in inputs.get('holdings',[])]
        for start in range(0,len(instruments),8):
            begin('PORTFOLIO INSTRUMENT SNAPSHOT',subtitle)
            rows=[]
            for item in instruments[start:start+8]:
                rows.append([str(item.get(k,'Unavailable')) for k in ['Symbol','Name','Sector','Price','Analyst Rating','Price Target Low','Price Target Average','Price Target High']])
            headers=['Ticker','Company','Sector','Current price','Analyst rating','Low target','Average target','High target']
            data=[[Paragraph(escape(str(v)),styles['MSCell']) for v in row] for row in [headers]+rows]
            t=Table(data,colWidths=[45,180,95,80,95,65,75,75]);t.setStyle(table_style(font_size=7))
            _,th=t.wrap(w-52,h);t.drawOn(c,26,h-85-th)
        # Split wide ledger into two readable tables; repeat dates for cash-flow reconciliation.
        table=pd.DataFrame(payload['table'])
        period='Date' if 'Date' in table else 'index' if 'index' in table else table.columns[0]
        table=table.rename(columns={period:'Date / period'});period='Date / period'
        groups=[[period,'Beginning Balance','Profit','Return %','Actual Withdrawal','Ending Balance'],
                [period,'Requested Withdrawal','Shortfall','Cumulative Withdrawals','Net Profit incl. Withdrawals']]
        if name=='Performance':
            table=performance_table(table)
            groups=[[period,'Beginning Balance','Profit','Return %','Ending Balance','Cumulative Profit']]
        for cols in groups:
            for start in range(0,len(table),18):
                begin(name.upper()+' YTD RESULTS',subtitle)
                data=[[Paragraph(escape(v),styles['MSCell']) for v in cols]]
                for _,row in table.iloc[start:start+18].iterrows():
                    data.append([str(row[v]) if v==period else (f"{row[v]:+.2f}%" if v=='Return %' else f"${row[v]:,.2f}") for v in cols])
                t=Table(data,colWidths=[(w-52)/len(cols)]*len(cols));t.setStyle(table_style(font_size=8))
                _,th=t.wrap(w-52,h);t.drawOn(c,26,h-85-th)
    c.save()
    return out.getvalue()


def build_ytd_excel(record):
    """Use the app's installed Excel engine for deployable runtime exports."""
    record=presentation_record(record)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import LineChart, Reference
    wb=Workbook(); summary=wb.active;summary.title='Summary'
    summary.append(['MarketScope YTD portfolio',record['name']])
    for k,v in record['inputs'].items():summary.append([k,', '.join(v) if isinstance(v,list) else v])
    summary.append(['Positive periods use pre-withdrawal returns; partial months included.'])
    summary.append(['Strategy','Start','Through','Coverage','Invested','Current balance','Profit incl withdrawals','Return','Positive days','Days','Positive months','Months','Withdrawn'])
    for name,payload in record['strategies'].items():
        s=metrics(payload)
        summary.append([name,payload['start_date'],payload['through'],coverage_label(payload),s['Beginning Balance'],s['Current Balance'],s['Profit / Loss'],s['Total Return %']/100,s['Positive Days'],s['Days'],s['Positive Months'],s['Months'],s['Withdrawn']])
        for j in [5,6,7,13]:summary.cell(summary.max_row,j).number_format='$#,##0.00'
        summary.cell(summary.max_row,8).number_format='0.00%'
        for kind in ['table','daily']:
            frame=pd.DataFrame(payload[kind]); sheet=wb.create_sheet(('Performance' if name=='Performance' else 'RB' if name=='Rebalanced' else 'NR')+' '+kind)
            date_col='Date' if 'Date' in frame else 'index' if 'index' in frame else frame.columns[0]
            frame=frame[[date_col]+[col for col in frame if col!=date_col]]
            if kind=='daily':frame['Cumulative profit']=frame['Ending Balance']+frame['Actual Withdrawal'].cumsum()-s['Beginning Balance']
            if name=='Performance':frame=performance_table(frame)
            sheet.append(frame.columns.tolist())
            for row in frame.itertuples(index=False,name=None):sheet.append(list(row))
            for j,col in enumerate(frame.columns,1):
                if j>1:
                    for cells in sheet.iter_rows(min_row=2,min_col=j,max_col=j):cells[0].number_format='0.00"%"' if col=='Return %' else '$#,##0.00'
            sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions
            if kind=='daily':
                chart=LineChart();chart.title='Cumulative profit / loss';chart.y_axis.title='USD';chart.x_axis.title='Trading date'
                chart.add_data(Reference(sheet,min_col=sheet.max_column,min_row=1,max_row=sheet.max_row),titles_from_data=True)
                chart.set_categories(Reference(sheet,min_col=1,min_row=2,max_row=sheet.max_row));sheet.add_chart(chart,'L2')
    if record.get('instruments'):
        sheet=wb.create_sheet('Instruments')
        frame=pd.DataFrame(record['instruments'])
        sheet.append(frame.columns.tolist())
        for row in frame.itertuples(index=False,name=None):sheet.append(list(row))
        sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions
    for sheet in wb:
        for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='123B40');cell.font=Font(color='56E58B',bold=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width=min(42,max(20,max(len(str(c.value or '')) for c in col)+2))
        for cell in sheet[1]:cell.alignment=Alignment(wrap_text=True)
    out=BytesIO();wb.save(out);return out.getvalue()


def get_exports(record):
    import streamlit as st
    @st.cache_data(show_spinner=False,max_entries=24)
    def exports(serialized):
        saved=json.loads(serialized)
        return build_ytd_pdf(saved),build_ytd_excel(saved)
    return exports(json.dumps(record,sort_keys=True))


def render_downloads(record, key, columns=None):
    import streamlit as st
    try:
        pdf,xlsx=get_exports(record)
        a,b=columns if columns is not None else st.columns(2)
        a.download_button('Download YTD PDF',pdf,'MarketScope_YTD.pdf','application/pdf',key=key+'_pdf')
        b.download_button('Download YTD Excel',xlsx,'MarketScope_YTD.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',key=key+'_xlsx')
    except Exception as exc:
        st.error(f'Report generation failed; simulation results remain available: {exc}')
