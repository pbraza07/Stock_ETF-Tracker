from io import BytesIO
import pandas as pd
import pytest
from pypdf import PdfReader
from openpyxl import load_workbook
from ytd_simulator import simulate_ytd, make_saved_record
from ytd_reports import build_ytd_pdf, build_ytd_excel, summary_html

def test_reports_preserve_partial_coverage_and_cashflows():
    p=pd.DataFrame({'A':[100,110,120]},index=pd.to_datetime(['2026-01-02','2026-01-05','2026-02-02']))
    r=make_saved_record({n:simulate_ytd(p,1000,'Daily',20,'Daily',rb,'2026-02-02') for n,rb in [('Rebalanced',True),('Non-Rebalanced',False)]},{'holdings':['A'],'cadence':'Daily','withdrawal':20},'Saved <YTD>')
    pdf=PdfReader(BytesIO(build_ytd_pdf(r)))
    text=' '.join(p.extract_text() for p in pdf.pages)
    assert 'Partial YTD 2026-01-02' in text
    assert 'CUMULATIVE PROFIT / LOSS' in text
    assert 'Actual Withdrawal' in text
    wb=load_workbook(BytesIO(build_ytd_excel(r)))
    assert {'Summary','RB table','NR table','RB daily','NR daily'} == set(wb.sheetnames)
    assert len(wb['RB daily']._charts)==1
    daily=r['strategies']['Rebalanced']['daily']
    assert wb['RB daily'].cell(3,9).value == pytest.approx(daily[-1]['Ending Balance'] + 40 - 1000)
    html=summary_html(r)
    assert 'Saved &lt;YTD&gt;' in html and 'REBALANCE DIFFERENCE' in html
