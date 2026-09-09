from io import BytesIO
from urllib.parse import urlparse,parse_qs
import pandas as pd
from pypdf import PdfReader
from openpyxl import load_workbook
from economic_calendar import calendar_url,calendar_embed
from ytd_simulator import simulate_ytd,make_saved_record
from ytd_reports import build_ytd_pdf,build_ytd_excel,summary_html,performance_table

def test_none_legacy_save_has_one_performance_export_with_dates():
    prices=pd.DataFrame({'A':[100,110,120]},index=pd.to_datetime(['2025-12-31','2026-01-02','2026-01-05']))
    output=simulate_ytd(prices,1000,'Daily',0,'None',as_of='2026-01-05')
    record=make_saved_record({'Rebalanced':output,'Non-Rebalanced':output},{'cadence':'None','holdings':['A']},'Example')
    text=' '.join(page.extract_text() for page in PdfReader(BytesIO(build_ytd_pdf(record))).pages)
    assert 'Rebalanced' not in text and 'REBALANCED' not in text
    assert all(date in text for date in ['2025-12-31','2026-01-02','2026-01-05'])
    assert 'Date / period' in text
    wb=load_workbook(BytesIO(build_ytd_excel(record)))
    assert 'Performance table' in wb.sheetnames and 'RB table' not in wb.sheetnames
    assert 'REBALANCE' not in summary_html(record)
    assert 'Actual Withdrawal' not in performance_table(output[0])

def test_calendar_us_only_and_full_frame_dark_filter():
    params=parse_qs(urlparse(calendar_url()).query)
    assert params['countries']==['5'] and params['importance']==['3']
    assert params['defaultFont']==['#111827']
    assert 'filter:invert(.94)' in calendar_embed()
